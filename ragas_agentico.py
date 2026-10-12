"""
ragas_agentico.py
=================
Roda o RAGAS (Faithfulness, Answer Relevancy e Context Relevance) sobre as respostas do
RAG Agêntico, com a mesma configuração do Artigo 1:
    RAGAS 0.2.x | avaliador GPT-3.5-turbo, temperatura 0 |
    embeddings paraphrase-multilingual-MiniLM-L12-v2 (para Answer Relevancy)

Uso (no MESMO ambiente Python em que você rodou o RAGAS do Artigo 1):

    python ragas_agentico.py --estimar resultados_agentico_formatado.xlsx resultados_agentico.csv
        -> NÃO chama nenhuma API: só diz quantas chamadas e quantos tokens devem ser gastos.

    python ragas_agentico.py --apenas Q1 Q8 resultados_agentico_formatado.xlsx resultados_agentico.csv
        -> roda só Q1 e Q8 (teste barato antes de rodar tudo).

    python ragas_agentico.py resultados_agentico_formatado.xlsx resultados_agentico.csv
        -> roda as 20 perguntas (repetição 1). Grava em ragas_agentico.csv.

Recursos de segurança (você está sem crédito sobrando):
  * Antes de começar, faz uma chamada mínima para checar chave, crédito e modelo.
  * Grava o resultado a cada pergunta; se o crédito acabar, para e retoma depois.
  * Se uma pergunta falhar nas 3 métricas, o script PARA (em vez de gastar chamadas à toa).
  * Bloqueia a mistura de configurações diferentes no mesmo arquivo de saída.

IMPORTANTE: mantenha a MESMA versão do ragas do Artigo 1 (pip show ragas). Não atualize
para 0.3 ou superior: os nomes e o cálculo das métricas mudam.

Arquivos de entrada: .csv ou .xlsx com as colunas id, pergunta, contextos (lista em JSON),
resposta e, opcionalmente, repeticao e erro (o formato gerado pelo benchmark).
"""

import argparse
import csv
import json
import math
import os
import sys

from dotenv import load_dotenv
from openpyxl import load_workbook

csv.field_size_limit(2**31 - 1)

# ---------- configuração do avaliador (igual ao Artigo 1) ----------
MODELO_AVALIADOR = "gpt-3.5-turbo"
TEMPERATURA_AVALIADOR = 0
MODELO_EMBEDDING = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
# O Context Relevance do ragas 0.2 junta todos os trechos em um só texto e lê apenas os
# primeiros 7000 caracteres. Contextos maiores que isso são truncados na avaliação.
LIMITE_CONTEXT_RELEVANCE = 7000

# Resultados do Artigo 1 (Tabela 1): (Faithfulness, Answer Relevancy, Context Relevance)
ARTIGO1 = {
    "Q1": (1.00, 0.73, 1.00), "Q2": (0.00, 0.81, 0.50), "Q3": (0.80, 0.57, 0.50),
    "Q4": (0.50, 0.52, 0.50), "Q5": (0.00, 0.78, 1.00), "Q6": (0.83, 0.63, 1.00),
    "Q7": (0.75, 0.62, 1.00), "Q8": (0.00, 0.81, 0.00), "Q9": (0.78, 0.82, 0.75),
    "Q10": (0.50, 0.72, 0.25), "Q11": (0.87, 0.66, 0.75), "Q12": (0.67, 0.75, 1.00),
    "Q13": (0.78, 0.56, 1.00), "Q14": (1.00, 0.81, 1.00), "Q15": (1.00, 0.81, 1.00),
    "Q16": (0.50, 0.65, 1.00), "Q17": (0.50, 0.69, 1.00), "Q18": (0.75, 0.56, 1.00),
    "Q19": (0.88, 0.69, 0.25), "Q20": (0.50, 0.86, 0.25),
}

METRICAS = ("faithfulness", "answer_relevancy", "context_relevance")
CAMPOS_SAIDA = ["id", "repeticao", *METRICAS, "n_chunks", "chars_contexto",
                "contexto_truncado", "completo"]


# ============================================================
# Leitura dos dados
# ============================================================
def ler_arquivo(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".xlsx":
        wb = load_workbook(caminho, read_only=True, data_only=True)
        ws = wb["resultados"] if "resultados" in wb.sheetnames else wb.worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c) for c in linhas[0]]
        return [{cab[j]: ("" if v is None else str(v)) for j, v in enumerate(l)} for l in linhas[1:]]
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def carregar_amostras(arquivos, apenas, todas_rep, incluir_f):
    amostras, vistos, ignoradas = [], set(), []
    for caminho in arquivos:
        linhas = [{str(k).strip().lower(): ("" if v is None else v) for k, v in l.items()}
                  for l in ler_arquivo(caminho)]
        if linhas:
            faltam = {"id", "pergunta", "contextos", "resposta"} - set(linhas[0])
            if faltam:
                sys.exit(
                    f"\nO arquivo '{caminho}' não tem as colunas {sorted(faltam)}.\n"
                    f"Colunas encontradas: {list(linhas[0])[:8]}\n"
                    "O RAGAS precisa dos trechos recuperados (coluna 'contextos'). Use os arquivos de "
                    "RESULTADOS do benchmark (resultados_agentico.csv e resultados_agentico_formatado.xlsx), "
                    "não a planilha de avaliação."
                )
        for n_linha, l in enumerate(linhas, start=2):
            origem = f"{os.path.basename(caminho)} (linha {n_linha}, id={l.get('id') or '?'})"
            try:
                qid = str(l["id"]).strip().upper()
                rep = int(l.get("repeticao") or 1)
                resposta = str(l.get("resposta", "")).strip()
                if l.get("erro") or not resposta:
                    ignoradas.append(origem + ": sem resposta ou com erro")
                    continue
                if not qid.startswith("Q") and not incluir_f:
                    continue
                if rep != 1 and not todas_rep:
                    continue
                if apenas and qid not in apenas:
                    continue
                if (qid, rep) in vistos:
                    sys.exit(f"Execução duplicada nos arquivos de entrada: {qid} repetição {rep}.")
                vistos.add((qid, rep))
                contextos = json.loads(l["contextos"]) if isinstance(l["contextos"], str) else l["contextos"]
                if not contextos:
                    raise ValueError("contextos vazios")
            except (ValueError, TypeError) as e:  # inclui json.JSONDecodeError
                ignoradas.append(f"{origem}: dados inválidos ({e})")
                continue
            amostras.append({"id": qid, "repeticao": rep, "pergunta": str(l["pergunta"]).strip(),
                             "contextos": contextos, "resposta": l["resposta"]})
    if ignoradas:
        print("AVISO: linhas ignoradas nos arquivos de entrada:")
        for x in ignoradas:
            print("  -", x)
    amostras.sort(key=lambda a: (int(a["id"][1:]) if a["id"][1:].isdigit() else 999, a["repeticao"]))
    return amostras


# ============================================================
# Estimativa (sem API)
# ============================================================
def estimar(amostras):
    n = len(amostras)
    tokens_in = tokens_out = 0
    for a in amostras:
        ctx = sum(len(c) for c in a["contextos"])
        r = len(a["resposta"])
        chars = 3 * r + ctx + 2 * min(ctx, LIMITE_CONTEXT_RELEVANCE)
        tokens_in += chars / 3.5 + 5 * 350          # texto + instruções dos prompts
        tokens_out += r / 3.5 + 300
    truncadas = [a["id"] for a in amostras if sum(len(c) for c in a["contextos"]) > LIMITE_CONTEXT_RELEVANCE]
    print(f"Amostras: {n}")
    print(f"Chamadas ao avaliador (aprox.): {n * 5}  "
          f"(Faithfulness ~2, Answer Relevancy ~1, Context Relevance 2 por amostra)")
    print(f"Tokens de entrada (aprox.): {tokens_in:,.0f} | tokens de saída (aprox.): {tokens_out:,.0f}")
    print("Estimativa grosseira. Multiplique pelos preços atuais do gpt-3.5-turbo no painel da OpenAI.")
    print(f"Amostras com contexto acima de {LIMITE_CONTEXT_RELEVANCE} caracteres "
          f"(truncado no Context Relevance): {truncadas or 'nenhuma'}")


# ============================================================
# Avaliação
# ============================================================
def exigir_ragas_02(ignorar=False):
    """O Artigo 1 usou o RAGAS 0.2.x. As métricas mudam entre versões, então rodar em outra
    versão invalidaria a comparação (e poderia gastar crédito à toa)."""
    try:
        import ragas
    except ImportError:
        sys.exit("O pacote ragas não está instalado neste ambiente. Veja as instruções de instalação "
                 "do ambiente com ragas==0.2.15.")
    if str(ragas.__version__).startswith("0.2") or ignorar:
        return
    sys.exit(
        f"\nEste ambiente tem o ragas {ragas.__version__}, mas o Artigo 1 usou a versão 0.2.x. "
        "As métricas mudam entre versões, e rodar assim invalida a comparação.\n"
        "Crie um ambiente separado com: pip install \"ragas==0.2.15\" (veja as instruções completas). "
        "Para rodar mesmo assim, use --ignorar-versao (e declare a versão no artigo)."
    )


def montar_avaliador():
    from langchain_openai import ChatOpenAI
    from ragas.embeddings import LangchainEmbeddingsWrapper
    from ragas.llms import LangchainLLMWrapper
    try:
        from langchain_huggingface import HuggingFaceEmbeddings
    except ImportError:
        from langchain_community.embeddings import HuggingFaceEmbeddings

    chat = ChatOpenAI(model=MODELO_AVALIADOR, temperature=TEMPERATURA_AVALIADOR)
    try:  # chamada mínima: confere chave, crédito e disponibilidade do modelo
        chat.invoke("Responda apenas: ok")
    except Exception as e:
        sys.exit(f"\nFalha ao chamar a OpenAI (chave inválida, sem crédito ou modelo indisponível):\n{e}")
    return (LangchainLLMWrapper(chat),
            LangchainEmbeddingsWrapper(HuggingFaceEmbeddings(model_name=MODELO_EMBEDDING)))


def avaliar_amostra(a, llm, emb):
    from ragas import EvaluationDataset, evaluate
    from ragas.metrics import AnswerRelevancy, ContextRelevance, Faithfulness
    from ragas.run_config import RunConfig

    ds = EvaluationDataset.from_list([{
        "user_input": a["pergunta"],
        "retrieved_contexts": a["contextos"],
        "response": a["resposta"],
    }])
    res = evaluate(ds, metrics=[Faithfulness(), AnswerRelevancy(), ContextRelevance()],
                   llm=llm, embeddings=emb, raise_exceptions=False, show_progress=False,
                   run_config=RunConfig(max_workers=4, timeout=120))
    s = res.scores[0]
    nan = float("nan")
    return {"faithfulness": s.get("faithfulness", nan),
            "answer_relevancy": s.get("answer_relevancy", nan),
            "context_relevance": s.get("nv_context_relevance", nan)}


def e_nan(x):
    return x is None or x == "" or (isinstance(x, float) and math.isnan(x))


def carregar_saida(caminho):
    if not os.path.exists(caminho):
        return {}
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        return {(r["id"], int(r["repeticao"])): r for r in csv.DictReader(f)}


def configuracao():
    try:
        import ragas
        versao = ragas.__version__
    except Exception:
        versao = "desconhecida"
    return {"avaliador": MODELO_AVALIADOR, "temperatura": TEMPERATURA_AVALIADOR,
            "embedding": MODELO_EMBEDDING, "ragas": versao}


def resumo(saida, caminho_cmp):
    completos = [r for r in saida.values() if str(r["completo"]) == "True" and int(r["repeticao"]) == 1]
    completos.sort(key=lambda r: int(r["id"][1:]))
    if not completos:
        print("Nenhuma amostra completa para resumir.")
        return
    print("\n" + "=" * 78)
    print(f"{'ID':5}| {'Faithfulness':^17}| {'Answer Relevancy':^17}| {'Context Relevance':^19}| trunc.")
    print(f"{'':5}| {'agêntico':>8} {'A1':>8}| {'agêntico':>8} {'A1':>8}| {'agêntico':>9} {'A1':>9}|")
    linhas_cmp, soma = [], {m: [0.0, 0.0, 0] for m in METRICAS}
    for r in completos:
        a1 = ARTIGO1.get(r["id"])
        vals = [float(r[m]) for m in METRICAS]
        cel = []
        for i, m in enumerate(METRICAS):
            ref = a1[i] if a1 else float("nan")
            cel.append(f"{vals[i]:8.2f} {ref:8.2f}")
            if a1:
                soma[m][0] += vals[i]; soma[m][1] += ref; soma[m][2] += 1
        print(f"{r['id']:5}| {cel[0]}| {cel[1]}| {cel[2]}  | {'sim' if str(r['contexto_truncado']) == 'True' else ''}")
        linhas_cmp.append({"id": r["id"], **{m: vals[i] for i, m in enumerate(METRICAS)},
                           **{f"{m}_artigo1": (a1[i] if a1 else "") for i, m in enumerate(METRICAS)},
                           **{f"{m}_diferenca": (round(vals[i] - a1[i], 3) if a1 else "") for i, m in enumerate(METRICAS)},
                           "n_chunks": r["n_chunks"], "contexto_truncado": r["contexto_truncado"]})
    print("-" * 78)
    medias = []
    for m in METRICAS:
        s_ag, s_a1, n = soma[m]
        medias.append(f"{s_ag / n:8.2f} {s_a1 / n:8.2f}" if n else " " * 17)
    print(f"{'Média':5}| {medias[0]}| {medias[1]}| {medias[2]}  | (n = {soma[METRICAS[0]][2]})")
    with open(caminho_cmp, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas_cmp[0].keys()))
        w.writeheader(); w.writerows(linhas_cmp)
    print(f"\nComparação salva em {caminho_cmp}")
    print("Obs.: 'sim' em trunc. = contexto > 7000 caracteres; o Context Relevance leu só o início.")


def main():
    load_dotenv()
    ap = argparse.ArgumentParser(description="RAGAS sobre as respostas do RAG Agêntico")
    ap.add_argument("arquivos", nargs="+", help="CSV/XLSX de resultados do benchmark")
    ap.add_argument("--estimar", action="store_true", help="só estima custo, sem chamar API")
    ap.add_argument("--apenas", nargs="*", default=[], help="IDs a avaliar (ex.: Q1 Q8)")
    ap.add_argument("--saida", default="ragas_agentico.csv")
    ap.add_argument("--todas-repeticoes", action="store_true")
    ap.add_argument("--incluir-f", action="store_true", help="inclui F1-F5 (o Artigo 1 não as avaliou com RAGAS)")
    ap.add_argument("--ignorar-versao", action="store_true", help="não exige o ragas 0.2.x")
    args = ap.parse_args()

    amostras = carregar_amostras(args.arquivos, {x.upper() for x in args.apenas},
                                 args.todas_repeticoes, args.incluir_f)
    if not amostras:
        sys.exit("Nenhuma amostra encontrada nos arquivos (confira ids, erros e filtros).")
    print(f"{len(amostras)} amostras: {[a['id'] for a in amostras]}")

    if args.estimar:
        estimar(amostras)
        return

    exigir_ragas_02(args.ignorar_versao)

    if not os.getenv("OPENAI_API_KEY"):
        sys.exit("OPENAI_API_KEY não encontrada. Coloque OPENAI_API_KEY=... no .env da pasta de execução.")

    cfg = configuracao()
    caminho_cfg = args.saida + ".config.json"
    if os.path.exists(caminho_cfg) and os.path.exists(args.saida):
        anterior = json.load(open(caminho_cfg, encoding="utf-8"))
        if anterior != cfg:
            sys.exit("A configuração do avaliador mudou desde a rodada anterior deste arquivo. "
                     "Use outro --saida ou apague o arquivo.\n"
                     f"Anterior: {anterior}\nAtual:    {cfg}")
    json.dump(cfg, open(caminho_cfg, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print("Configuração do avaliador:", cfg)

    saida = carregar_saida(args.saida)

    def salvar():
        with open(args.saida, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=CAMPOS_SAIDA)
            w.writeheader(); w.writerows(saida.values())

    pendentes = [a for a in amostras if str(saida.get((a["id"], a["repeticao"]), {}).get("completo")) != "True"]
    print(f"Já completas: {len(amostras) - len(pendentes)} | a avaliar agora: {len(pendentes)}")

    if pendentes:
        llm, emb = montar_avaliador()
        for a in pendentes:
            print(f"Avaliando {a['id']} (repetição {a['repeticao']})...")
            s = avaliar_amostra(a, llm, emb)
            if all(e_nan(s[m]) for m in METRICAS):
                salvar()
                sys.exit(f"\nAs três métricas falharam em {a['id']}. Provável falta de crédito ou erro de "
                         "chave/modelo. Parei para não gastar mais chamadas; o progresso foi salvo.")
            ctx = sum(len(c) for c in a["contextos"])
            saida[(a["id"], a["repeticao"])] = {
                "id": a["id"], "repeticao": a["repeticao"],
                **{m: ("" if e_nan(s[m]) else round(float(s[m]), 4)) for m in METRICAS},
                "n_chunks": len(a["contextos"]), "chars_contexto": ctx,
                "contexto_truncado": ctx > LIMITE_CONTEXT_RELEVANCE,
                "completo": not any(e_nan(s[m]) for m in METRICAS),
            }
            salvar()  # grava a cada amostra
            falhas = [m for m in METRICAS if e_nan(s[m])]
            if falhas:
                print(f"  aviso: métricas sem valor em {a['id']}: {falhas} (será refeita na próxima rodada)")

    resumo(saida, args.saida.replace(".csv", "") + "_comparacao.csv")


if __name__ == "__main__":
    main()